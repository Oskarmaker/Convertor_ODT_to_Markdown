import re
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches
import markdown2
import requests
import os
from PIL import Image
import io
import tempfile
from styles import Styles


def markdown_to_docx(markdown_content, output_path, image_base_path='.'):
    doc = Document()
    styles = Styles(doc) # Подгрузка стилей
    styles.run()

    def preprocess_markdown(content): # Поиск странных(нетипичных) указаний на картинки
        pattern_image = r'!\[\[(.*?)\]\]'

        def replace_image(match):
            filename = match.group(1).strip()
            full_path = os.path.join(image_base_path, filename)
            return f'![{filename}]({full_path})'

        content = re.sub(pattern_image, replace_image, content)

        pattern_nested_image = r'\[!([^]]*)\]\(([^)]+)\)'

        def replace_nested_image(match):
            alt_text = match.group(1).strip()
            url = match.group(2).strip()
            return f'![{alt_text}]({url})'

        content = re.sub(pattern_nested_image, replace_nested_image, content)
        return content

    markdown_content = preprocess_markdown(markdown_content)
    html_content = markdown2.markdown(markdown_content, extras=['tables', 'fenced-code-blocks'])
    soup = BeautifulSoup(html_content, 'html.parser')
    elements = soup.body.children if soup.body else soup.children

    def convert_to_png(image_path): # Конвертация в .png
        try:
            img = Image.open(image_path)
            if img.format.lower() != 'png':
                output = io.BytesIO()
                img.save(output, format='PNG')
                return output.getvalue()
            return open(image_path, 'rb').read()
        except Exception as e:
            print(f"Error converting image {image_path}: {e}")
            return None

    def process_element(element):
        if not element.name:
            return
        # Заголовки
        if element.name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            level = min(int(element.name[1]), 3)
            heading = doc.add_heading(element.get_text(), level=level)
            heading.style = doc.styles[f'Heading {level}']
        # Основной текст
        elif element.name == 'p':
            paragraph = doc.add_paragraph(style='Normal')
            for child in element.children:
                if child.name:
                    run = paragraph.add_run()
                    process_inline_elements(child, run)
                elif child.string and child.string.strip():
                    paragraph.add_run(child.string)
        # Списки
        elif element.name in ['ul', 'ol']:
            for li in element.find_all('li', recursive=False):
                paragraph = doc.add_paragraph(style='List Bullet' if element.name == 'ul' else 'List Number')
                for child in li.children:
                    if child.name:
                        run = paragraph.add_run()
                        process_inline_elements(child, run)
                    elif child.string and child.string.strip():
                        paragraph.add_run(child.string)
        # Таблицы
        elif element.name == 'table':
            rows = element.find_all('tr')
            if rows:
                cols_count = max(len(row.find_all(['td', 'th'])) for row in rows)
                table = doc.add_table(rows=len(rows), cols=cols_count)
                table.style = 'Table Grid'
                for row_idx, row in enumerate(rows):
                    cells = row.find_all(['td', 'th'])
                    for col_idx, cell in enumerate(cells):
                        cell_paragraph = table.rows[row_idx].cells[col_idx].paragraphs[0]
                        cell_paragraph.style = doc.styles['Текст таблицы']
                        for child in cell.children:
                            if child.name:
                                run = cell_paragraph.add_run()
                                process_inline_elements(child, run)
                            elif child.string and child.string.strip():
                                cell_paragraph.add_run(child.string)
                        if cell.name == 'th':
                            for run in cell_paragraph.runs:
                                run.bold = True
        # Картинки
        elif element.name == 'img':
            src = element.get('src')
            if src:
                print(f"Attempting to add image with src: {src}")
                paragraph = doc.add_paragraph(style='Рисунок')
                try:
                    if src.startswith('http'):
                        response = requests.get(src, timeout=10)
                        response.raise_for_status()
                        img_data = response.content
                        with tempfile.NamedTemporaryFile(delete=False, suffix='.png') as tmp:
                            img = Image.open(io.BytesIO(img_data))
                            img.save(tmp.name, 'PNG')
                            paragraph.add_run().add_picture(tmp.name, width=Inches(3))
                            os.unlink(tmp.name)
                    elif src[src.rfind('.'):] == '.gif':
                        paragraph.add_run().add_picture(src, width=Inches(3))
                    else:
                        img_data = convert_to_png(src)
                        if img_data:
                            with tempfile.NamedTemporaryFile(delete=False, suffix='.png') as tmp:
                                tmp.write(img_data)
                                tmp.flush()
                                paragraph.add_run().add_picture(tmp.name, width=Inches(3))
                                os.unlink(tmp.name)
                        else:
                            print(f"Image conversion failed: {src}")
                            paragraph.add_run(f"[Image not found or failed to convert: {src}]")
                except Exception as e:
                    print(f"Error adding image {src}: {e}")
                    # paragraph.add_run(f"[Error loading image: {src}]")

        for child in element.children:
            if child.name:
                process_element(child)
    # Выделение текста
    def process_inline_elements(element, run):
        if element.name == 'strong':
            run.bold = True
        elif element.name in ['em', 'i']:
            run.italic = True
        elif element.name == 'u':
            run.underline = True
        if element.string:
            run.text += element.string
        for child in element.children:
            if child.name:
                process_inline_elements(child, run)

    for element in elements:
        process_element(element)

    doc.save(output_path)
    print(f"Document saved to {output_path}")


path = 'C:/Users/vovat/Desktop/АКМС/гайд по настройки На русском.md' # Путь к файлу .md
with open(path, 'r', encoding='utf-8') as f:
    line = f.read()
markdown_content = line # Текст файла .md
path_to_docx = path[:path.rfind('.')] + '.docx' # Путь к файлу .docx
image_path = 'C:/Users/vovat/Desktop/АКМС/' # Папка с картинками


markdown_to_docx(markdown_content, path_to_docx, image_path     )